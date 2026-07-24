"""Convert the mid-term Markdown draft to a styled .docx.

Handles: H1/H2/H3 headings, markdown tables, bullet lists, ordered lists,
blockquotes, and inline bold (**text**) + inline code (`text`).
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

# 默认转换脚本同目录下的草稿；也可通过命令行参数指定源/目标文件
_BASE = Path(__file__).resolve().parent
SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else _BASE / "项目中期文档_草稿.md"
DST = Path(sys.argv[2]) if len(sys.argv) > 2 else _BASE / "项目中期文档.docx"

INLINE_RE = re.compile(r"(\*\*.+?\*\*|`.+?`)")


def add_inline(paragraph, text: str) -> None:
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.color.rgb = RGBColor(0x8A, 0x1A, 0x6B)
        else:
            paragraph.add_run(part)


def flush_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    ncol = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=ncol)
    table.style = "Light Grid Accent 1"
    for r, row in enumerate(rows):
        for c in range(ncol):
            cell = table.cell(r, c)
            cell.text = ""
            para = cell.paragraphs[0]
            add_inline(para, row[c] if c < len(row) else "")
            if r == 0:
                for run in para.runs:
                    run.bold = True
    doc.add_paragraph()


def parse_table_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def main() -> None:
    doc = Document()
    doc.styles["Normal"].font.name = "Microsoft YaHei"
    doc.styles["Normal"].font.size = Pt(10.5)

    lines = SRC.read_text(encoding="utf-8").splitlines()
    i = 0
    table_buf: list[list[str]] = []

    def flush():
        nonlocal table_buf
        if table_buf:
            flush_table(doc, table_buf)
            table_buf = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # table detection
        if stripped.startswith("|") and stripped.endswith("|"):
            # skip separator rows like |---|---|
            if re.fullmatch(r"\|[\s:\-|]+\|", stripped):
                i += 1
                continue
            table_buf.append(parse_table_row(stripped))
            i += 1
            continue
        else:
            flush()

        if not stripped:
            i += 1
            continue

        if stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("# "):
            h = doc.add_heading(stripped[2:], level=0)
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif stripped.startswith("> "):
            p = doc.add_paragraph(style="Intense Quote")
            add_inline(p, stripped[2:])
        elif re.match(r"^[-*] ", stripped):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, stripped[2:])
        elif re.match(r"^\d+\. ", stripped):
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\d+\. ", "", stripped))
        else:
            p = doc.add_paragraph()
            add_inline(p, stripped)
        i += 1

    flush()
    doc.save(DST)
    print(f"Saved: {DST}")


if __name__ == "__main__":
    main()
